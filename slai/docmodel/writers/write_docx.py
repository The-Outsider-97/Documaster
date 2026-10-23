"""DOCX writer (python-docx). Real Word structure: heading/list/quote styles, numbered lists that restart, tables with
merged cells and repeating header rows, hyperlinks, inline images with alt text, page setup, per-section headers and footers
with PAGE/NUMPAGES fields, page breaks, document properties.

Footnotes are written as a numbered "Notes" list at the end (python-docx has no footnote API); the result says so.
"""
from __future__ import annotations

import copy
import io
import re
from pathlib import Path
from typing import Any

from ...errors import DependencyMissing
from ..ir import has_page_break, heading_level, span_segments
from . import WriteOptions, WriteResult
from .common import (Notes, PageSetup, asset_bytes, char_style_table, clean, group_blocks, page_setup,
                     style_table, table_grid)

MONO = "Consolas"


def write_docx(package: Any, dest: Path, opts: WriteOptions) -> WriteResult:
    try:
        import docx
        from docx.enum.section import WD_ORIENT, WD_SECTION
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
    except ImportError as exc:
        raise DependencyMissing("Writing DOCX needs the python-docx package", details={"package": "python-docx"}) from exc

    doc_model = package.document
    notes = Notes.of(doc_model)
    styles_para = style_table(doc_model)
    styles_char = char_style_table(doc_model)
    sections = {s["id"]: s for s in doc_model.get("sections", [])}
    section_order = [s["id"] for s in doc_model.get("sections", [])]
    document = docx.Document()
    warnings: list[str] = []

    base_font = opts.font_family or "Calibri"
    normal = document.styles["Normal"]
    normal.font.name = base_font
    normal.font.size = Pt(opts.font_size_pt)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), base_font)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), doc_model.get("language", "en") if doc_model.get("language") != "und" else "en-US")
    rpr.append(lang)

    core = document.core_properties
    core.title = clean(opts.title or doc_model["title"])[:250]
    authors = (doc_model.get("publication") or {}).get("authors") or ([opts.author] if opts.author else [])
    if authors:
        core.author = clean("; ".join(authors))[:250]
    core.language = doc_model.get("language", "") if doc_model.get("language") != "und" else ""
    core.comments = "Created with SLAI Documaster"

    # ---- page setup helpers
    def apply_page(section: Any, setup: PageSetup) -> None:
        w, h = setup.width_mm, setup.height_mm
        section.orientation = WD_ORIENT.LANDSCAPE if w > h else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Mm(w), Mm(h)
        section.top_margin, section.bottom_margin = Mm(setup.top), Mm(setup.bottom)
        section.left_margin, section.right_margin = Mm(setup.left), Mm(setup.right)
        cols = section._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            section._sectPr.append(cols)
        cols.set(qn("w:num"), str(max(1, setup.columns)))
        cols.set(qn("w:space"), str(int(setup.column_gap_mm * 56.7)))

    def text_width_mm(setup: PageSetup) -> float:
        return max(20.0, setup.width_mm - setup.left - setup.right)

    def effective_setup(section: dict[str, Any] | None) -> PageSetup:
        setup = page_setup(section)
        if opts.page_size and not (section and "page" in section):
            from .common import PAGE_SIZES_MM
            if opts.page_size in PAGE_SIZES_MM:
                setup.width_mm, setup.height_mm = PAGE_SIZES_MM[opts.page_size]
        if opts.margins_mm is not None and not (section and "page" in section):
            setup.top = setup.right = setup.bottom = setup.left = opts.margins_mm
        return setup

    # ---- fields and header/footer
    def add_field(paragraph: Any, instr: str) -> None:
        for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
            run = paragraph.add_run()
            if kind:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                run._r.append(fc)
            elif text == instr:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = f" {instr} "
                run._r.append(it)
            else:
                run.text = text or ""

    def fill_template(paragraph: Any, template: str) -> None:
        template = template.replace("{{", "\x01").replace("}}", "\x02")
        for token in re.split(r"(\{page\}|\{pages\}|\{title\})", template):
            if token == "{page}":
                add_field(paragraph, "PAGE")
            elif token == "{pages}":
                add_field(paragraph, "NUMPAGES")
            elif token == "{title}":
                paragraph.add_run(clean(opts.title or doc_model["title"]))
            elif token:
                paragraph.add_run(clean(token.replace("\x01", "{").replace("\x02", "}")))

    def apply_header_footer(docx_section: Any, model_section: dict[str, Any] | None, setup: PageSetup) -> None:
        for edge, part in (("header", docx_section.header), ("footer", docx_section.footer)):
            template = (model_section or {}).get(edge)
            if template is None and edge == "footer" and opts.page_numbers and model_section is None:
                template = {"center": "{page}"}
            part.is_linked_to_previous = False
            paragraph = part.paragraphs[0]
            paragraph.style = document.styles["Normal"]      # the built-in Header style carries its own tab stops
            for r in list(paragraph.runs):
                r._r.getparent().remove(r._r)
            if not template:
                continue
            width = Mm(text_width_mm(setup))
            paragraph.paragraph_format.tab_stops.add_tab_stop(int(width / 2), WD_TAB_ALIGNMENT.CENTER)
            paragraph.paragraph_format.tab_stops.add_tab_stop(int(width), WD_TAB_ALIGNMENT.RIGHT)
            fill_template(paragraph, template.get("left", ""))
            paragraph.add_run("\t")
            fill_template(paragraph, template.get("center", ""))
            paragraph.add_run("\t")
            fill_template(paragraph, template.get("right", ""))
            for r in paragraph.runs:
                r.font.size = Pt(9)

    first_model = sections.get(section_order[0]) if section_order else None
    setup = effective_setup(first_model)
    apply_page(document.sections[0], setup)
    apply_header_footer(document.sections[0], first_model, setup)
    current_section_id = section_order[0] if section_order else None

    # ---- list numbering
    numbering_root = None
    try:
        numbering_root = document.part.numbering_part.element
    except Exception:  # noqa: BLE001
        numbering_root = None

    def new_num_for_style(style_name: str) -> int | None:
        if numbering_root is None:
            return None
        style = document.styles[style_name]
        ppr = style.element.pPr
        if ppr is None or ppr.numPr is None or ppr.numPr.numId is None:
            return None
        base_id = str(ppr.numPr.numId.val)
        abstract = None
        for num in numbering_root.findall(qn("w:num")):
            if num.get(qn("w:numId")) == base_id:
                abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        if abstract is None:
            return None
        ids = [int(n.get(qn("w:numId"))) for n in numbering_root.findall(qn("w:num"))]
        new_id = max(ids + [0]) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), abstract)
        num.append(a)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        numbering_root.append(num)
        return new_id

    # ---- inline content
    def hyperlink(paragraph: Any, url: str, text: str, props: dict[str, Any]) -> None:
        rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rid)
        run = paragraph.add_run(text)
        style_run(run, props)
        run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        run.font.underline = True
        link.append(run._r)
        paragraph._p.append(link)

    def style_run(run: Any, attrs: dict[str, Any], base: dict[str, Any] | None = None) -> None:
        merged = {**(base or {}), **attrs}
        if merged.get("bold"):
            run.bold = True
        if merged.get("italic"):
            run.italic = True
        if merged.get("underline"):
            run.underline = True
        if merged.get("strike"):
            run.font.strike = True
        if merged.get("superscript"):
            run.font.superscript = True
        if merged.get("subscript"):
            run.font.subscript = True
        if merged.get("small_caps"):
            run.font.small_caps = True
        if merged.get("code"):
            run.font.name = MONO
            r = run._r.get_or_add_rPr().find(qn("w:rFonts"))
            if r is not None:
                for a in ("w:ascii", "w:hAnsi", "w:cs"):
                    r.set(qn(a), MONO)
        if merged.get("font_family") and not merged.get("code"):
            run.font.name = merged["font_family"]
        if merged.get("font_size_pt"):
            run.font.size = Pt(merged["font_size_pt"])
        if merged.get("color"):
            run.font.color.rgb = RGBColor.from_string(merged["color"][1:].upper())
        if merged.get("background"):
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), merged["background"][1:].upper())
            run._r.get_or_add_rPr().append(shd)

    def add_inline(paragraph: Any, block: dict[str, Any], base: dict[str, Any] | None = None) -> None:
        for seg, attrs in span_segments(block["text"], block.get("spans")):
            props = dict(attrs)
            if props.get("style") in styles_char:
                props = {**styles_char[props["style"]], **props}
            text = clean(seg)
            if props.get("href"):
                hyperlink(paragraph, props["href"], text, {**(base or {}), **props})
            else:
                run = paragraph.add_run(text)
                style_run(run, props, base)
            if props.get("ref"):
                kind, number = notes.use(props["ref"])
                if number:
                    marker = paragraph.add_run(f"[{number}]")
                    marker.font.superscript = True

    def apply_paragraph_style(paragraph: Any, block: dict[str, Any]) -> dict[str, Any]:
        base: dict[str, Any] = {}
        props = styles_para.get(block.get("style", ""), {})
        pf = paragraph.paragraph_format
        align = (block.get("attributes") or {}).get("align") or props.get("align")
        if align:
            paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                                   "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
        if "line_height" in props:
            pf.line_spacing = props["line_height"]
        if "space_before_pt" in props:
            pf.space_before = Pt(props["space_before_pt"])
        if "space_after_pt" in props:
            pf.space_after = Pt(props["space_after_pt"])
        if "indent_mm" in props:
            pf.left_indent = Mm(props["indent_mm"])
        if "first_line_indent_mm" in props:
            pf.first_line_indent = Mm(props["first_line_indent_mm"])
        if props.get("keep_with_next"):
            pf.keep_with_next = True
        for key in ("bold", "italic", "underline", "strike", "small_caps", "font_family", "font_size_pt", "color", "background"):
            if key in props:
                base[key] = props[key]
        if has_page_break(block):
            pf.page_break_before = True
        return base

    def style_exists(name: str) -> bool:
        try:
            document.styles[name]
            return True
        except KeyError:
            return False

    body_width_mm = [text_width_mm(setup)]

    def maybe_new_section(block: dict[str, Any]) -> None:
        nonlocal current_section_id
        sid = block.get("section")
        if not sid or sid == current_section_id:
            return
        model = sections.get(sid)
        new_setup = effective_setup(model)
        docx_section = document.add_section(WD_SECTION.NEW_PAGE)
        apply_page(docx_section, new_setup)
        apply_header_footer(docx_section, model, new_setup)
        body_width_mm[0] = text_width_mm(new_setup)
        current_section_id = sid

    # ---- blocks
    for kind, item in group_blocks(doc_model["blocks"]):
        if kind == "list":
            run_blocks = item
            nums: dict[int, int | None] = {}
            for block in run_blocks:
                maybe_new_section(block)
                attrs = block.get("attributes") or {}
                depth = min(int(attrs.get("depth", 0)), 8)
                ordered = bool(attrs.get("ordered"))
                for d in [d for d in nums if d > depth]:
                    del nums[d]
                level_name = ("List Number" if ordered else "List Bullet") + ("" if depth == 0 else f" {min(depth + 1, 3)}")
                style_name = level_name if style_exists(level_name) else ("List Number" if ordered else "List Bullet")
                p = document.add_paragraph(style=style_name)
                if ordered:
                    if depth not in nums:
                        nums[depth] = new_num_for_style(style_name)
                    if nums[depth] is not None:
                        ppr = p._p.get_or_add_pPr()
                        num_pr = OxmlElement("w:numPr")
                        ilvl = OxmlElement("w:ilvl")
                        ilvl.set(qn("w:val"), "0")
                        nid = OxmlElement("w:numId")
                        nid.set(qn("w:val"), str(nums[depth]))
                        num_pr.append(ilvl)
                        num_pr.append(nid)
                        ppr.append(num_pr)
                else:
                    nums.pop(depth, None)
                if depth >= 3:
                    p.paragraph_format.left_indent = Mm(6.35 * (depth + 1))
                base = apply_paragraph_style(p, block)
                add_inline(p, block, base)
            continue
        block = item
        maybe_new_section(block)
        t = block["type"]
        if t == "heading":
            level = heading_level(block)
            style_name = f"Heading {level}" if style_exists(f"Heading {level}") else "Normal"
            p = document.add_paragraph(style=style_name)
            base = apply_paragraph_style(p, block)
            add_inline(p, block, base)
        elif t == "quote":
            p = document.add_paragraph(style="Quote" if style_exists("Quote") else "Normal")
            base = apply_paragraph_style(p, block)
            add_inline(p, block, base)
        elif t == "code":
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            ppr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F2F2F2")
            ppr.append(shd)
            if has_page_break(block):
                p.paragraph_format.page_break_before = True
            run = p.add_run(clean(block["text"]))
            style_run(run, {"code": True})
            run.font.size = Pt(max(8, opts.font_size_pt - 1.5))
        elif t == "table":
            if block["text"]:
                cap = document.add_paragraph(style="Caption" if style_exists("Caption") else "Normal")
                cap.add_run(clean(block["text"]))
                cap.paragraph_format.keep_with_next = True
                if has_page_break(block):
                    cap.paragraph_format.page_break_before = True
            _docx_table(document, block, style_exists, qn, OxmlElement, WD_TABLE_ALIGNMENT, WD_ALIGN_PARAGRAPH, Pt)
            document.add_paragraph()
        elif t == "figure":
            _docx_figure(document, package, block, body_width_mm[0], style_exists, qn, Mm, WD_ALIGN_PARAGRAPH, warnings)
        else:
            if not block["text"].strip():
                if has_page_break(block):  # an empty anchor (a page without extractable text) keeps its page
                    p = document.add_paragraph()
                    p.paragraph_format.page_break_before = True
                continue
            p = document.add_paragraph()
            base = apply_paragraph_style(p, block)
            add_inline(p, block, base)

    if opts.include_notes and (notes.note_numbers or notes.cited):
        h = document.add_paragraph(style="Heading 2" if style_exists("Heading 2") else "Normal")
        h.add_run("Notes")
        for n, text in notes.footnotes():
            p = document.add_paragraph()
            r = p.add_run(f"[{n}] ")
            r.font.superscript = True
            p.add_run(clean(text))
        for line in notes.citation_lines():
            document.add_paragraph(clean(line))
        warnings.append("Footnotes are written as a numbered 'Notes' list at the end (python-docx cannot create native footnotes)")
    document.save(str(dest))
    return WriteResult(dest, "docx", "python-docx", warnings=warnings,
                       kept=["headings (Heading 1-6)", "lists with restarting numbers", "tables with merged cells and header rows",
                             "bold/italic/underline/strike/super/subscript/colour", "hyperlinks", "images with alt text",
                             "page size, margins, columns", "headers and footers with page numbers", "page breaks",
                             "document title, author, language"],
                       lost=["native footnotes (written as a Notes list)", "comments and revision history", "signatures"])


def _docx_table(document: Any, block: dict[str, Any], style_exists: Any, qn: Any, OxmlElement: Any,
                WD_TABLE_ALIGNMENT: Any, WD_ALIGN_PARAGRAPH: Any, Pt: Any) -> None:
    grid = table_grid(block["table"])
    table = document.add_table(rows=len(grid.rows), cols=grid.columns)
    if style_exists("Table Grid"):
        table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    origins = []
    for r, row in enumerate(grid.rows):
        for c, cell in enumerate(row):
            if cell is None:
                continue
            origins.append((r, c, cell))
            target = table.cell(r, c)
            target.text = ""
            para = target.paragraphs[0]
            run = para.add_run(clean(cell["text"]))
            if r < grid.header_rows or cell.get("header"):
                run.bold = True
            align = cell.get("align")
            if align:
                para.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                                  "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    for r, c, cell in origins:
        cs, rs = cell.get("colspan", 1), cell.get("rowspan", 1)
        if cs > 1 or rs > 1:
            table.cell(r, c).merge(table.cell(r + rs - 1, c + cs - 1))
    for r in range(grid.header_rows):
        tr = table.rows[r]._tr
        trpr = tr.get_or_add_trPr()
        flag = OxmlElement("w:tblHeader")
        trpr.append(flag)


def _docx_figure(document: Any, package: Any, block: dict[str, Any], width_mm: float, style_exists: Any, qn: Any,
                 Mm: Any, WD_ALIGN_PARAGRAPH: Any, warnings: list[str]) -> None:
    fig = block["figure"]
    got = asset_bytes(package, fig["asset"])
    if not got:
        warnings.append(f"Image {fig['asset']} is missing from the package and was skipped")
        return
    data, _ = got
    paragraph = document.add_paragraph()
    if has_page_break(block):
        paragraph.paragraph_format.page_break_before = True
    align = fig.get("align", "center")
    paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                           "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    target = min(float(fig.get("width_mm", width_mm)), width_mm)
    run = paragraph.add_run()
    picture = run.add_picture(io.BytesIO(data), width=Mm(target))
    doc_pr = picture._inline.docPr
    alt = fig.get("alt", "")
    if alt:
        doc_pr.set("descr", clean(alt))
    if block["text"]:
        cap = document.add_paragraph(style="Caption" if style_exists("Caption") else "Normal")
        cap.add_run(clean(block["text"]))
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
