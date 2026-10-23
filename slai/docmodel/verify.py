"""Verification of every file Documaster writes.

A written file is checked three ways, and the result says exactly which checks ran:

1. SLAI Reader ``IntegrityEngine.verify_output`` (where Reader knows the format)
2. the file is re-opened with a real parser for its format (pypdf, python-docx, Documaster's ODT reader, ElementTree, ...)
3. the words of the re-opened file are compared with the words of the document model that was written: *coverage* (how many of
   the intended words are present) and *order* (how much of the sequence is intact).

Layout, fonts and appearance are not judged here; the report does not claim they are.
"""
from __future__ import annotations

import difflib
import re
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from .. import textstats
from ..bridge import sldc
from ..services import Services

READER_FORMAT = {"pdf": "pdf", "docx": "docx", "html": "html", "md": "md", "txt": "txt", "xml": "xml"}
COVERAGE_MIN = 0.995


def words_of(text: str) -> list[str]:
    return textstats.tokenize(unicodedata.normalize("NFKC", text))


def rtf_to_text(raw: str) -> str:
    """Small RTF text extractor used only for verification (handles \\uN, \\'hh, groups to skip, \\par/\\line/\\tab)."""
    out: list[str] = []
    skip_depth = None
    depth = 0
    i, n = 0, len(raw)
    skip_dest = ("fonttbl", "colortbl", "stylesheet", "listtable", "listoverridetable", "info", "pict", "generator", "fldinst",
                 "listtext", "pntext", "picprop", "header", "footer")
    pending_skip = 0
    while i < n:
        ch = raw[i]
        if ch == "{":
            depth += 1
            if raw.startswith("{\\*", i):
                if skip_depth is None:
                    skip_depth = depth
            i += 1
        elif ch == "}":
            if skip_depth is not None and depth == skip_depth:
                skip_depth = None
            depth -= 1
            i += 1
        elif ch == "\\":
            m = re.match(r"\\([a-zA-Z]+)(-?\d+)? ?", raw[i:i + 40])
            if m:
                word, arg = m.group(1), m.group(2)
                i += m.end()
                if word in skip_dest and skip_depth is None:
                    skip_depth = depth
                elif skip_depth is None:
                    if word in ("par", "line", "row", "cell"):
                        out.append("\n")
                    elif word == "tab":
                        out.append("\t")
                    elif word == "u" and arg is not None:
                        code = int(arg)
                        out.append(chr(code + 65536 if code < 0 else code))
                        if i < n and raw[i] == "?":
                            i += 1
            elif i + 1 < n and raw[i + 1] == "'":
                if skip_depth is None:
                    out.append(chr(int(raw[i + 2:i + 4], 16)))
                i += 4
            else:
                if skip_depth is None and i + 1 < n and raw[i + 1] in "\\{}":
                    out.append(raw[i + 1])
                i += 2
        else:
            if skip_depth is None and ch not in "\r\n":
                out.append(ch)
            i += 1
    text = "".join(out)
    # UTF-16 surrogate pairs written as two \u escapes
    return text.encode("utf-16", "surrogatepass").decode("utf-16", "replace")


def extract_output_text(path: Path, fmt: str, svc: Services, password: str | None = None) -> tuple[str, dict[str, Any]]:
    """Re-open ``path`` with an independent parser; returns (text, facts)."""
    facts: dict[str, Any] = {}
    if fmt == "pdf":
        import pypdf
        reader = pypdf.PdfReader(str(path))
        if reader.is_encrypted:
            reader.decrypt(password or "")
        facts["pages"] = len(reader.pages)
        return "\n".join((p.extract_text() or "") for p in reader.pages), facts
    if fmt == "docx":
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        d = docx.Document(str(path))
        parts: list[str] = []
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        for child in d.element.body.iterchildren():      # document order: paragraphs and tables interleaved
            if child.tag == ns + "p":
                parts.append(Paragraph(child, d).text)
            elif child.tag == ns + "tbl":
                for tr in Table(child, d)._tbl.tr_lst:
                    for tc in tr.tc_lst:
                        parts.append("\n".join(p.text for p in docx.table._Cell(tc, d).paragraphs))
        facts["paragraphs"] = len(d.paragraphs)
        return "\n".join(parts), facts
    if fmt == "odt":
        from .read_odt import read_odt
        data = read_odt(str(path))
        pkg = sldc().new_package("v", data["blocks"], assets=data["assets"])
        return sldc().extract_text(pkg, validate=False), facts
    if fmt == "html":
        from .read_html import read_html_bytes
        data = read_html_bytes(path.read_bytes())
        pkg = sldc().new_package("v", data["blocks"], assets=data["assets"])
        return sldc().extract_text(pkg, validate=False), facts
    if fmt == "rtf":
        raw = path.read_bytes().decode("latin-1")
        facts["balancedBraces"] = _balanced(raw)
        if not raw.startswith("{\\rtf1"):
            raise ValueError("not an RTF header")
        return rtf_to_text(raw), facts
    if fmt == "xml":
        root = ET.parse(str(path)).getroot()
        skip = {"data"}
        parts = [t for el in root.iter() if el.tag.split("}")[-1] not in skip for t in (el.text or "", el.tail or "")]
        return " ".join(parts), facts
    if fmt == "sldc":
        pkg = sldc().read_package(path, password=password)
        facts["revision"] = sldc().revision_id(pkg)
        return sldc().extract_text(pkg, validate=False), facts
    return path.read_text(encoding="utf-8"), facts


def _balanced(raw: str) -> bool:
    depth = 0
    escaped = False
    for ch in raw:
        if escaped:
            escaped = False
        elif ch == "\\":
            escaped = True
        elif ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth < 0:
                return False
    return depth == 0


def compare_words(expected: list[str], actual: list[str]) -> dict[str, Any]:
    if not expected:
        return {"coverage": 1.0, "order": 1.0, "expected": 0, "found": 0, "missingSample": [], "orderChecked": True}
    exp, act = Counter(expected), Counter(actual)
    found = sum(min(c, act.get(w, 0)) for w, c in exp.items())
    missing = [w for w, c in exp.items() if act.get(w, 0) < c][:10]
    result: dict[str, Any] = {"coverage": round(found / len(expected), 5), "expected": len(expected), "found": found,
                              "missingSample": missing, "orderChecked": False, "order": None}
    if len(expected) <= 40_000 and len(actual) <= 120_000:
        matcher = difflib.SequenceMatcher(None, expected, actual, autojunk=False)
        matched = sum(b.size for b in matcher.get_matching_blocks())
        result["order"] = round(matched / len(expected), 5)
        result["orderChecked"] = True
    return result


def verify_written(path: Path, fmt: str, package: Any, svc: Services, *, password: str | None = None,
                   expected_text: str | None = None, check_text: bool = True) -> dict[str, Any]:
    """Run all checks on a freshly written file. Never raises: failures are reported in the result."""
    checks: list[dict[str, Any]] = []
    facts: dict[str, Any] = {}
    reader_check = {"ran": False, "valid": None, "failed": [], "detail": "Reader does not check this format"}
    if fmt in READER_FORMAT:
        reader_check = svc.reader.verify_output(path, READER_FORMAT[fmt])
    checks.append({"name": "SLAI Reader integrity check", "ok": bool(reader_check["valid"]) if reader_check["ran"] else None,
                   "ran": reader_check["ran"],
                   "detail": ("failed: " + ", ".join(reader_check["failed"])) if reader_check["ran"] and not reader_check["valid"]
                   else (reader_check["detail"] if not reader_check["ran"] else "passed")})
    if fmt == "sldc":       # Reader has no SLDC check; SLAI's own SLDC validator is the authority for this format
        try:
            pkg = sldc().read_package(path, password=password)
            sldc().validate_package(pkg)
            checks.append({"name": "SLAI SLDC validator", "ok": True, "ran": True, "detail": "the package is valid SLDC"})
        except Exception as exc:  # noqa: BLE001
            checks.append({"name": "SLAI SLDC validator", "ok": False, "ran": True, "detail": f"{type(exc).__name__}: {str(exc)[:200]}"})
    text = ""
    try:
        text, facts = extract_output_text(path, fmt, svc, password)
        checks.append({"name": "File re-opened with an independent parser", "ok": True, "ran": True, "detail": "opened"})
    except Exception as exc:  # noqa: BLE001
        checks.append({"name": "File re-opened with an independent parser", "ok": False, "ran": True,
                       "detail": f"{type(exc).__name__}: {str(exc)[:200]}"})
    comparison: dict[str, Any] | None = None
    if check_text and checks[-1]["ok"]:
        expected = words_of(expected_text if expected_text is not None else sldc().extract_text(package, validate=False))
        comparison = compare_words(expected, words_of(text))
        ok = comparison["coverage"] >= COVERAGE_MIN
        detail = (f"{comparison['found']} of {comparison['expected']} words present ({comparison['coverage'] * 100:.2f}%)"
                  + (f"; order intact {comparison['order'] * 100:.1f}%" if comparison.get("order") is not None else "; order not checked (large)"))
        if comparison["missingSample"]:
            detail += "; e.g. missing: " + ", ".join(comparison["missingSample"][:5])
        checks.append({"name": "Text matches the document model", "ok": ok, "ran": True, "detail": detail})
    ran = [c for c in checks if c["ran"]]
    return {"ok": all(c["ok"] for c in ran if c["ok"] is not None), "checks": checks, "facts": facts,
            "comparison": comparison, "readerRan": reader_check["ran"]}
